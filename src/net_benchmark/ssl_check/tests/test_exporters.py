"""Tests for `net_benchmark.ssl_check.exporters`.

Two things this file is specifically watching for, because both were real
bugs during development:

1. openpyxl reads embedded chart images LAZILY, at workbook.save() time, not
   when add_image() is called — deleting the temp PNGs before save() raises
   deep inside the writer. The chart-cleanup-order test exists because of
   exactly that failure.
2. A withheld percentile (item 57) must render as a blank cell, never 0.00 —
   a 0.00 in a spreadsheet reads as a measurement.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Callable, Coroutine, Dict, List, Tuple

import pytest
from openpyxl import load_workbook

try:
    import weasyprint  # noqa: F401

    _WEASYPRINT_AVAILABLE = True
except ImportError:
    _WEASYPRINT_AVAILABLE = False

from net_benchmark.ssl_check.analysis import SSLAnalyzer, parse_threshold
from net_benchmark.ssl_check.core import (
    PolicyConfig,
    SSLCheckEngine,
    SSLResult,
    SSLTarget,
    evaluate_policy,
)
from net_benchmark.ssl_check.exporters import (
    SSLCSVExporter,
    SSLExcelExporter,
    SSLExportBundle,
    SSLPDFExporter,
    build_provenance,
)
from net_benchmark.ssl_check.handshake import HandshakeStatus, StartTLSProtocol

from .conftest import TLSServerHandle

TLSServerFactory = Callable[..., Coroutine[None, None, TLSServerHandle]]


@pytest.fixture
async def small_fleet(
    tls_server: TLSServerFactory, unreachable_target: Tuple[str, int]
) -> List[SSLResult]:
    healthy = await tls_server(validity_days=150)
    expiring = await tls_server(validity_days=90, age_days=87, filename_hint="exp")
    engine = SSLCheckEngine(
        handshake_samples=6,
        min_samples=5,
        warmup_handshakes=0,
        connect_timeout=3,
        handshake_timeout=6,
    )
    unreachable_host, unreachable_port = unreachable_target
    results = await engine.check_targets(
        [
            SSLTarget("localhost", healthy.port, pinned_ip="127.0.0.1"),
            SSLTarget("localhost", expiring.port, pinned_ip="127.0.0.1"),
            SSLTarget(unreachable_host, unreachable_port),
        ]
    )
    policy = PolicyConfig(min_days_remaining=30)
    for result in results:
        evaluate_policy(result, policy)
    return results


class TestProvenance:
    """Item 56 — a multi-store trust verdict is unreproducible without this."""

    def test_provenance_has_store_and_version_fields(self) -> None:
        provenance = build_provenance()
        assert provenance["trust_store"] == "certifi"
        assert "trust_store_version" in provenance
        assert "openssl_version" in provenance
        assert provenance["openssl_version"] is not None

    def test_provenance_extra_merged(self) -> None:
        provenance = build_provenance(extra={"run_id": "abc123"})
        assert provenance["run_id"] == "abc123"


class TestCSVExport:
    async def test_raw_export_has_one_row_per_result(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        out = tmp_path / "raw.csv"
        SSLCSVExporter.export_raw_results(small_fleet, str(out))
        with open(out) as handle:
            rows = list(csv.DictReader(handle))
        assert len(rows) == 3

    async def test_certificate_flattened_not_json_blob(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        """A raw export whose interesting content is a JSON blob in one cell
        is not a CSV anyone can filter — the certificate must be flattened
        into cert_* columns."""
        out = tmp_path / "raw.csv"
        SSLCSVExporter.export_raw_results(small_fleet, str(out))
        with open(out) as handle:
            header = next(csv.reader(handle))
        assert "cert_issuer_cn" in header
        assert "cert_days_remaining" in header
        assert "cipher_iana_hex" in header

    async def test_summary_export_no_dict_columns(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        analyzer = SSLAnalyzer(small_fleet)
        out = tmp_path / "summary.csv"
        SSLCSVExporter.export_summary_statistics(analyzer, str(out))
        with open(out) as handle:
            header = next(csv.reader(handle))
        assert "status_counts" not in header
        assert "policy_failure_counts" not in header
        assert "expiry_alert_counts" not in header

    async def test_expiry_timeline_export_has_unknown_bucket(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        out = tmp_path / "timeline.csv"
        SSLCSVExporter.export_expiry_timeline(small_fleet, str(out))
        with open(out) as handle:
            rows = list(csv.DictReader(handle))
        buckets = {row["bucket"] for row in rows}
        assert "unknown" in buckets


class TestJSONExport:
    async def test_schema_and_provenance_present(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        analyzer = SSLAnalyzer(small_fleet)
        out = tmp_path / "results.json"
        SSLExportBundle.export_json(small_fleet, analyzer, str(out))
        payload = json.loads(out.read_text())
        assert payload["schema_version"] == 1
        assert payload["provenance"]["trust_store"] == "certifi"
        assert len(payload["results"]) == 3

    async def test_threshold_results_included_when_provided(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        analyzer = SSLAnalyzer(small_fleet)
        # get_thresholds_report() is the real per-target entry point the CLI
        # actually calls — see analysis.py's docstring on why thresholds are
        # evaluated per target rather than as a flat sequence.
        report = analyzer.get_thresholds_report([parse_threshold("total_checks>=0")])
        out = tmp_path / "results.json"
        SSLExportBundle.export_json(
            small_fleet, analyzer, str(out), threshold_results=report
        )
        payload = json.loads(out.read_text())
        assert "thresholds" in payload
        # One entry per (target, threshold) — every target in the fleet gets
        # its own row, each carrying which target it belongs to.
        assert len(payload["thresholds"]) == len(report)
        assert all("target" in entry for entry in payload["thresholds"])

    async def test_expiry_timeline_in_json(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        analyzer = SSLAnalyzer(small_fleet)
        out = tmp_path / "results.json"
        SSLExportBundle.export_json(small_fleet, analyzer, str(out))
        payload = json.loads(out.read_text())
        labels = {g["label"] for g in payload["expiry_timeline"]}
        assert "unknown" in labels


class TestExcelExport:
    async def test_all_expected_sheets_present(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        analyzer = SSLAnalyzer(small_fleet)
        out = tmp_path / "results.xlsx"
        SSLExcelExporter.export_results(small_fleet, analyzer, str(out))
        workbook = load_workbook(out)
        assert set(workbook.sheetnames) >= {
            "Summary",
            "Expiry Timeline",
            "Certificates",
            "Raw Results",
            "Provenance",
        }

    async def test_charts_sheet_saves_without_error(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        """Regression test for the specific bug this exporter had during
        development: openpyxl resolves embedded image files lazily at
        workbook.save() time. Deleting the temp chart PNGs before save()
        raises FileNotFoundError from deep inside the writer. If cleanup
        happens in the wrong order, THIS test fails with that traceback."""
        analyzer = SSLAnalyzer(small_fleet)
        out = tmp_path / "results.xlsx"
        SSLExcelExporter.export_results(
            small_fleet, analyzer, str(out), include_charts=True
        )
        assert out.exists()
        assert out.stat().st_size > 5000
        workbook = load_workbook(out)
        assert "Charts" in workbook.sheetnames

    async def test_temp_chart_files_cleaned_up(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        analyzer = SSLAnalyzer(small_fleet)
        out = tmp_path / "results.xlsx"
        import glob
        import tempfile

        before = set(glob.glob(f"{tempfile.gettempdir()}/**/ssl_*.png", recursive=True))
        SSLExcelExporter.export_results(small_fleet, analyzer, str(out))
        after = set(glob.glob(f"{tempfile.gettempdir()}/**/ssl_*.png", recursive=True))
        assert after - before == set(), "chart temp files were not cleaned up"

    async def test_withheld_percentile_is_blank_not_zero(
        self, tls_server: TLSServerFactory, tmp_path: Path
    ) -> None:
        """Item 57 — a withheld p95 rendered as 0.00 would read as a real
        measurement; it must be blank in the sheet."""
        handle = await tls_server()
        engine = SSLCheckEngine(
            handshake_samples=2,
            min_samples=5,
            warmup_handshakes=0,
            connect_timeout=3,
            handshake_timeout=6,
        )
        result = await engine.check_target(
            SSLTarget("localhost", handle.port, pinned_ip="127.0.0.1")
        )
        analyzer = SSLAnalyzer([result])
        out = tmp_path / "withheld.xlsx"
        SSLExcelExporter.export_results([result], analyzer, str(out))

        workbook = load_workbook(out)
        sheet = workbook["Summary"]
        header = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
        p95_index = header.index("Handshake p95 ms")
        first_data_row = [
            cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))
        ]
        assert first_data_row[p95_index] is None

    async def test_expiry_timeline_sheet_sorted_soonest_first(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        analyzer = SSLAnalyzer(small_fleet)
        out = tmp_path / "results.xlsx"
        SSLExcelExporter.export_results(small_fleet, analyzer, str(out))
        sheet = load_workbook(out)["Expiry Timeline"]
        rows = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2)]
        # Column 1 (index 1) is "Days remaining". The unreachable target has
        # None there and must sort last, not first as an implicit-zero sort
        # would place it.
        assert rows[-1][1] is None
        assert rows[-1][2] == "unknown"

    async def test_threshold_sheet_present_when_provided(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        analyzer = SSLAnalyzer(small_fleet)
        report = analyzer.get_thresholds_report([parse_threshold("total_checks>=0")])
        out = tmp_path / "results.xlsx"
        SSLExcelExporter.export_results(
            small_fleet, analyzer, str(out), threshold_results=report
        )
        workbook = load_workbook(out)
        assert "Thresholds" in workbook.sheetnames
        sheet = workbook["Thresholds"]
        header = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
        assert header[0] == "Target"

    async def test_no_charts_option_skips_charts_sheet(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        analyzer = SSLAnalyzer(small_fleet)
        out = tmp_path / "no_charts.xlsx"
        SSLExcelExporter.export_results(
            small_fleet, analyzer, str(out), include_charts=False
        )
        assert "Charts" not in load_workbook(out).sheetnames


class TestExportThresholdResultsCSV:
    """`SSLExportBundle.export_threshold_results()` — the per-target
    threshold CSV, called directly by the CLI's --threshold gate. Every row
    must carry which target it belongs to; a threshold CSV without that
    column is not useful once the console output that produced it is gone.
    """

    async def test_one_row_per_target_and_threshold(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        analyzer = SSLAnalyzer(small_fleet)
        report = analyzer.get_thresholds_report(
            [parse_threshold("total_checks>=0"), parse_threshold("success_rate>=0")]
        )
        path = SSLExportBundle.export_threshold_results(report, str(tmp_path), "run")
        with open(path) as handle:
            rows = list(csv.DictReader(handle))
        # 2 thresholds x however many targets get_thresholds_report produced.
        assert len(rows) == 2 * len(report)

    async def test_returns_the_written_path(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        analyzer = SSLAnalyzer(small_fleet)
        report = analyzer.get_thresholds_report([parse_threshold("total_checks>=0")])
        path = SSLExportBundle.export_threshold_results(report, str(tmp_path), "myrun")
        assert path == str(tmp_path / "myrun_thresholds.csv")
        assert Path(path).exists()

    async def test_every_row_carries_its_target(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        analyzer = SSLAnalyzer(small_fleet)
        report = analyzer.get_thresholds_report([parse_threshold("total_checks>=0")])
        path = SSLExportBundle.export_threshold_results(report, str(tmp_path), "run")
        with open(path) as handle:
            rows = list(csv.DictReader(handle))
        targets_in_csv = {row["target"] for row in rows}
        assert targets_in_csv == set(report.keys())
        assert "" not in targets_in_csv

    async def test_pass_and_fail_both_written_correctly(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        analyzer = SSLAnalyzer(small_fleet)
        # A threshold every target passes, and one every target fails, so
        # the CSV genuinely contains both outcomes rather than only one.
        report = analyzer.get_thresholds_report(
            [
                parse_threshold("total_checks>=0"),
                parse_threshold("total_checks>999999"),
            ]
        )
        path = SSLExportBundle.export_threshold_results(report, str(tmp_path), "run")
        with open(path) as handle:
            rows = list(csv.DictReader(handle))
        passed_values = {row["passed"] for row in rows}
        assert passed_values == {"True", "False"}

    def test_empty_report_still_writes_a_header_only_file(self, tmp_path: Path) -> None:
        path = SSLExportBundle.export_threshold_results({}, str(tmp_path), "empty")
        with open(path) as handle:
            rows = list(csv.DictReader(handle))
        assert rows == []
        assert Path(path).exists()


class TestSSLPDFExporter:
    """Item 52. `_generate_html()` is tested as the pure function it is;
    `export_results()` is tested against the REAL weasyprint installed in
    this dev environment, producing an actual PDF — not mocked — plus the
    missing-extra path with the import genuinely forced to fail."""

    def test_generate_html_includes_target_and_alert_class(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        analyzer = SSLAnalyzer(small_fleet)
        html = SSLPDFExporter._generate_html(small_fleet, analyzer, build_provenance())
        for result in small_fleet:
            assert result.target in html
        # worst_expiry_alert values double as CSS classes in the report —
        # confirm at least one real alert value made it into the markup.
        stats = analyzer.get_target_statistics()
        assert any(f"class='{s.worst_expiry_alert}'" in html for s in stats)

    def test_generate_html_includes_provenance(
        self, small_fleet: List[SSLResult]
    ) -> None:
        analyzer = SSLAnalyzer(small_fleet)
        provenance = build_provenance(extra={"run_id": "abc123-unique"})
        html = SSLPDFExporter._generate_html(small_fleet, analyzer, provenance)
        assert provenance["trust_store"] in html
        assert provenance["openssl_version"] in html

    async def test_unreachable_section_present_when_failures_exist(
        self, small_fleet: List[SSLResult]
    ) -> None:
        analyzer = SSLAnalyzer(small_fleet)
        html = SSLPDFExporter._generate_html(small_fleet, analyzer, build_provenance())
        failed = analyzer.get_failed_targets()
        assert len(failed) > 0, "fixture must include an unreachable target"
        assert "Unreachable" in html
        assert failed[0][0] in html

    def test_unreachable_section_absent_when_no_failures(self) -> None:
        """The conditional expression that adds the whole Unreachable
        section must actually be conditional — confirmed by giving it a
        fleet with nothing unreachable."""

        class _AllHealthyAnalyzer:
            def get_overall_statistics(self) -> Dict[str, object]:
                return {
                    "measured_checks": 1,
                    "total_checks": 1,
                    "hosts": 1,
                    "certificates_observed": 1,
                    "cert_expiry_days_min": 90,
                }

            def get_target_statistics(self) -> List[object]:
                return []

            def get_failed_targets(self) -> List[object]:
                return []

        html = SSLPDFExporter._generate_html(
            [], _AllHealthyAnalyzer(), build_provenance()  # type: ignore[arg-type]
        )
        assert "Unreachable" not in html

    @pytest.mark.skipif(not _WEASYPRINT_AVAILABLE, reason="weasyprint not installed")
    async def test_export_results_produces_a_real_pdf(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        """Not mocked: this actually invokes weasyprint and writes a real
        PDF file, verified by its magic bytes rather than trusting that
        write_pdf() was merely called."""
        analyzer = SSLAnalyzer(small_fleet)
        out = tmp_path / "report.pdf"
        SSLPDFExporter.export_results(small_fleet, analyzer, str(out))
        assert out.exists()
        content = out.read_bytes()
        assert content[:5] == b"%PDF-", "not a real PDF file"
        assert len(content) > 500, "suspiciously small for a rendered report"

    async def test_export_results_raises_runtime_error_without_weasyprint(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        """Forces the actual ImportError weasyprint's absence would raise,
        rather than asserting on the source code — sys.modules['weasyprint']
        = None makes the import statement itself fail exactly as it would
        if the [pdf] extra were never installed."""
        import sys
        from unittest.mock import patch

        analyzer = SSLAnalyzer(small_fleet)
        out = tmp_path / "report.pdf"
        with patch.dict(sys.modules, {"weasyprint": None}):
            with pytest.raises(RuntimeError, match=r"\[pdf\] extra"):
                SSLPDFExporter.export_results(small_fleet, analyzer, str(out))
        assert not out.exists()


# ---------------------------------------------------------------------------
# Per-host grade (0.6.1 items 19-20)
# ---------------------------------------------------------------------------


def _base_result(target: str = "example.com", port: int = 443) -> SSLResult:
    return SSLResult(
        host=target,
        port=port,
        starttls=StartTLSProtocol.NONE,
        status=HandshakeStatus.OK,
        start_time=0.0,
        end_time=0.0,
        measured=True,
    )


class TestHostGrade:
    def test_nothing_attempted_is_none(self) -> None:
        from net_benchmark.ssl_check.exporters import _compute_host_grades

        grades = _compute_host_grades([_base_result()])
        assert len(grades) == 1
        assert grades[0].overall_ok is None

    def test_chain_verified_true_is_ok(self) -> None:
        from net_benchmark.ssl_check.chain import ChainAudit
        from net_benchmark.ssl_check.exporters import _compute_host_grades

        result = _base_result()
        result.chain_audit = ChainAudit(attempted=True, verified=True)
        grades = _compute_host_grades([result])
        assert grades[0].overall_ok is True
        assert grades[0].chain_verified is True

    def test_chain_verified_false_is_fail(self) -> None:
        from net_benchmark.ssl_check.chain import ChainAudit
        from net_benchmark.ssl_check.exporters import _compute_host_grades

        result = _base_result()
        result.chain_audit = ChainAudit(
            attempted=True, verified=False, verification_error="untrusted root"
        )
        grades = _compute_host_grades([result])
        assert grades[0].overall_ok is False
        assert grades[0].chain_error == "untrusted root"

    def test_revoked_true_is_fail(self) -> None:
        from net_benchmark.ssl_check.exporters import _compute_host_grades
        from net_benchmark.ssl_check.revocation import CRLStatus, RevocationAudit

        result = _base_result()
        result.revocation_audit = RevocationAudit(
            attempted=True, crl_status=CRLStatus.REVOKED
        )
        grades = _compute_host_grades([result])
        assert grades[0].revoked is True
        assert grades[0].overall_ok is False

    def test_ct_untrusted_is_fail(self) -> None:
        from net_benchmark.ssl_check.ct import CTAudit
        from net_benchmark.ssl_check.exporters import _compute_host_grades

        result = _base_result()
        result.ct_audit = CTAudit(attempted=True, sct_trust=[])
        # all_trusted is None with no SCTs -- simulate an explicit untrusted
        # verdict instead via a fake attribute-compatible object is overkill;
        # exercise via the real property using a minimal SCTTrust stand-in.
        from datetime import datetime, timezone

        from net_benchmark.ssl_check.certificate import (
            SignedCertificateTimestampInfo,
        )
        from net_benchmark.ssl_check.ct import SCTTrust

        sct = SignedCertificateTimestampInfo(
            log_id_hex="00" * 32,
            timestamp=datetime.now(timezone.utc),
            version="v1",
            entry_type="PRE_CERTIFICATE",
            signature_algorithm="ECDSA",
        )
        result.ct_audit = CTAudit(
            attempted=True, sct_trust=[SCTTrust(sct=sct, log=None, trusted=False)]
        )
        grades = _compute_host_grades([result])
        assert grades[0].ct_all_trusted is False
        assert grades[0].overall_ok is False

    def test_lint_warning_only_is_ok(self) -> None:
        from net_benchmark.ssl_check.exporters import _compute_host_grades
        from net_benchmark.ssl_check.lint import (
            FindingSeverity,
            LintAudit,
            LintAvailability,
            LintFinding,
        )

        result = _base_result()
        result.lint_audit = LintAudit(
            attempted=True,
            availability=LintAvailability.AVAILABLE,
            findings=[
                LintFinding(
                    severity=FindingSeverity.WARNING,
                    code="x",
                    message=None,
                    node_path="cert",
                )
            ],
        )
        grades = _compute_host_grades([result])
        assert grades[0].lint_worst_severity == "warning"
        assert grades[0].overall_ok is True

    def test_lint_error_is_fail(self) -> None:
        from net_benchmark.ssl_check.exporters import _compute_host_grades
        from net_benchmark.ssl_check.lint import (
            FindingSeverity,
            LintAudit,
            LintAvailability,
            LintFinding,
        )

        result = _base_result()
        result.lint_audit = LintAudit(
            attempted=True,
            availability=LintAvailability.AVAILABLE,
            findings=[
                LintFinding(
                    severity=FindingSeverity.ERROR,
                    code="x",
                    message=None,
                    node_path="cert",
                )
            ],
        )
        grades = _compute_host_grades([result])
        assert grades[0].overall_ok is False

    def test_worst_case_across_multiple_results_for_same_target(self) -> None:
        """One result says the chain verified, another (a later sample)
        says it didn't -- the target's grade reports the failure, not
        whichever result happened to be processed first."""
        from net_benchmark.ssl_check.chain import ChainAudit
        from net_benchmark.ssl_check.exporters import _compute_host_grades

        good = _base_result()
        good.chain_audit = ChainAudit(attempted=True, verified=True)
        bad = _base_result()
        bad.chain_audit = ChainAudit(attempted=True, verified=False)

        grades = _compute_host_grades([good, bad])
        assert len(grades) == 1  # same target, one grade
        assert grades[0].chain_verified is False

    def test_weakest_cipher_across_results(self) -> None:
        from net_benchmark.ssl_check.enumeration import (
            CipherStrength,
            EnumerationResult,
        )
        from net_benchmark.ssl_check.exporters import _compute_host_grades

        strong = _base_result()
        strong.enumeration = EnumerationResult(attempted=True, ciphers=[])
        # weakest_supported_cipher_strength is a computed property from
        # `ciphers`; simplest reliable way to get a fixed value here is a
        # monkeypatch-free construction via a supported cipher list.
        from net_benchmark.ssl_check.enumeration import CipherSupport

        strong.enumeration.ciphers = [
            CipherSupport("A-SUITE", None, True, CipherStrength.A, "strong")
        ]
        weak = _base_result()
        weak.enumeration = EnumerationResult(
            attempted=True,
            ciphers=[CipherSupport("C-SUITE", None, True, CipherStrength.C, "weak")],
        )
        grades = _compute_host_grades([strong, weak])
        assert grades[0].weakest_cipher_strength == "C"

    def test_different_targets_produce_separate_grades(self) -> None:
        from net_benchmark.ssl_check.exporters import _compute_host_grades

        a = _base_result(target="a.example.com")
        b = _base_result(target="b.example.com")
        grades = _compute_host_grades([a, b])
        assert {g.target for g in grades} == {"a.example.com:443", "b.example.com:443"}

    def test_policy_failures_included_in_overall(self) -> None:
        from net_benchmark.ssl_check.exporters import _compute_host_grades

        result = _base_result()
        result.policy_failures = ["deprecated TLS version"]
        grades = _compute_host_grades([result])
        assert grades[0].overall_ok is False
        assert grades[0].policy_failures == ["deprecated TLS version"]


class TestPDFFindingsSection:
    def test_findings_section_present_when_checks_ran(
        self, small_fleet: List[SSLResult]
    ) -> None:
        from net_benchmark.ssl_check.chain import ChainAudit

        # small_fleet's results have no chain_audit by default; attach one
        # so the findings section has something to render.
        small_fleet[0].chain_audit = ChainAudit(attempted=True, verified=True)
        analyzer = SSLAnalyzer(small_fleet)
        html = SSLPDFExporter._generate_html(small_fleet, analyzer, build_provenance())
        assert "<h2>Findings summary</h2>" in html
        assert "verified" in html

    def test_no_findings_section_when_nothing_attempted(
        self, small_fleet: List[SSLResult]
    ) -> None:
        analyzer = SSLAnalyzer(small_fleet)
        html = SSLPDFExporter._generate_html(small_fleet, analyzer, build_provenance())
        assert "<h2>Findings summary</h2>" not in html


class TestExcelGradeSheet:
    def test_grade_sheet_present(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        from net_benchmark.ssl_check.chain import ChainAudit

        small_fleet[0].chain_audit = ChainAudit(
            attempted=True, verified=False, verification_error="test error"
        )
        analyzer = SSLAnalyzer(small_fleet)
        out = tmp_path / "report.xlsx"
        SSLExcelExporter.export_results(small_fleet, analyzer, str(out))
        wb = load_workbook(str(out))
        assert "Per-Host Grade" in wb.sheetnames
        ws = wb["Per-Host Grade"]
        rows = list(ws.iter_rows(values_only=True))
        assert rows[0][0] == "Target"
        # Find the row for the target we attached a failing chain to.
        target_col = rows[0].index("Target")
        overall_col = rows[0].index("Overall")
        matching = [r for r in rows[1:] if r[target_col] == small_fleet[0].target]
        assert matching
        assert matching[0][overall_col] == "FAIL"


class TestReportBranding:
    """SaaS-side-only report customization — no CLI flag exposes this,
    tested here as the importable engine capability it is. `None`
    (every call site the CLI itself uses) must be provably identical to
    omitting the parameter entirely, not just "close enough."
    """

    def test_pdf_unbranded_default_unchanged(
        self, small_fleet: List[SSLResult]
    ) -> None:
        analyzer = SSLAnalyzer(small_fleet)
        provenance = build_provenance()
        html_omitted = SSLPDFExporter._generate_html(small_fleet, analyzer, provenance)
        html_explicit_none = SSLPDFExporter._generate_html(
            small_fleet, analyzer, provenance, branding=None
        )
        assert html_omitted == html_explicit_none
        assert "<title>SSL/TLS Report</title>" in html_omitted
        assert "data:image" not in html_omitted

    def test_pdf_branded_title_and_organization(
        self, small_fleet: List[SSLResult]
    ) -> None:
        from net_benchmark.ssl_check.exporters import ReportBranding

        analyzer = SSLAnalyzer(small_fleet)
        branding = ReportBranding(
            report_title="Acme Corp TLS Report", organization_name="Acme Corporation"
        )
        html = SSLPDFExporter._generate_html(
            small_fleet, analyzer, build_provenance(), branding
        )
        assert "<title>Acme Corp TLS Report</title>" in html
        assert "<h1>Acme Corp TLS Report</h1>" in html
        assert "Acme Corporation" in html
        assert (
            "SSL/TLS Report" not in html.split("<body>")[1]
            if "<body>" in html
            else True
        )

    def test_pdf_branded_logo_embedded_as_base64(
        self, small_fleet: List[SSLResult]
    ) -> None:
        from net_benchmark.ssl_check.exporters import ReportBranding

        analyzer = SSLAnalyzer(small_fleet)
        branding = ReportBranding(logo_bytes=b"not-a-real-png-but-bytes-are-bytes")
        html = SSLPDFExporter._generate_html(
            small_fleet, analyzer, build_provenance(), branding
        )
        assert "data:image/png;base64," in html

    def test_excel_unbranded_default_unchanged(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        analyzer = SSLAnalyzer(small_fleet)
        out = tmp_path / "unbranded.xlsx"
        SSLExcelExporter.export_results(
            small_fleet, analyzer, str(out), include_charts=False
        )
        wb = load_workbook(str(out))
        assert wb["Summary"]["A1"].value != "Acme Corp TLS Report"
        # Default title falls back to whatever openpyxl itself sets when
        # never assigned -- confirm this code path never touched it.
        assert wb.properties.title in (None, "")

    def test_excel_branded_title_creator_and_logo(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        import io

        from PIL import Image as PILImage

        from net_benchmark.ssl_check.exporters import ReportBranding

        buf = io.BytesIO()
        PILImage.new("RGB", (4, 4), (10, 10, 10)).save(buf, format="PNG")

        analyzer = SSLAnalyzer(small_fleet)
        branding = ReportBranding(
            report_title="Acme Corp TLS Report",
            organization_name="Acme Corporation",
            logo_bytes=buf.getvalue(),
        )
        out = tmp_path / "branded.xlsx"
        SSLExcelExporter.export_results(
            small_fleet, analyzer, str(out), include_charts=False, branding=branding
        )
        wb = load_workbook(str(out))
        assert wb.properties.title == "Acme Corp TLS Report"
        assert wb.properties.creator == "Acme Corporation"
        sheet = wb["Summary"]
        assert sheet["A1"].value == "Acme Corp TLS Report"
        assert sheet["A2"].value == "Acme Corporation"
        assert len(sheet._images) == 1

    def test_excel_branding_without_logo_skips_image(
        self, small_fleet: List[SSLResult], tmp_path: Path
    ) -> None:
        from net_benchmark.ssl_check.exporters import ReportBranding

        analyzer = SSLAnalyzer(small_fleet)
        branding = ReportBranding(report_title="Title Only")
        out = tmp_path / "title_only.xlsx"
        SSLExcelExporter.export_results(
            small_fleet, analyzer, str(out), include_charts=False, branding=branding
        )
        wb = load_workbook(str(out))
        assert len(wb["Summary"]._images) == 0